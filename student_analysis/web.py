from __future__ import annotations

from io import BytesIO
import sys
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from flask import Flask, flash, redirect, render_template, request, send_file, url_for
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import Workbook
from werkzeug.utils import secure_filename

from .analysis import (
    class_names,
    class_summary,
    exam_names,
    exam_order_labels,
    filter_by_class,
    progress_matrix,
    student_exam_coverage,
    student_names,
    student_series,
)
from .charts import combined_trend_chart_png, combined_trend_chart_svg
from .importers import parse_excel_with_report
from .models import DEFAULT_LINES, ScoreLine, ScoreRecord, coerce_score, competition_ranks
from .ocr import image_to_text
from .storage import (
    append_scores,
    delete_preview,
    load_exam_order,
    load_lines,
    load_preview,
    load_scores,
    save_exam_order,
    save_lines,
    save_preview,
    save_scores,
)


UPLOAD_DIR = Path("uploads")


def create_app() -> Flask:
    base_dir = resource_path()
    app = Flask(
        __name__,
        instance_relative_config=False,
        template_folder=str(base_dir / "templates"),
        static_folder=str(base_dir / "static"),
    )
    app.config["SECRET_KEY"] = "local-dev-secret"
    UPLOAD_DIR.mkdir(exist_ok=True)

    @app.get("/")
    def index():
        records = load_scores()
        lines = load_lines(DEFAULT_LINES)
        manual_order = load_exam_order()
        classes = class_names(records)
        selected_class = request.args.get("class_name") or (classes[0] if classes else "")
        scoped_records = filter_by_class(records, selected_class)
        scoped_students = student_names(scoped_records)
        selected_student = request.args.get("student") or (scoped_students[0] if scoped_students else "")
        progress_sort = request.args.get("progress_sort", "score_delta_desc")
        series = student_series(scoped_records, selected_student, selected_class, manual_order) if selected_student else []
        ordered_exams = exam_order_labels(scoped_records, manual_order)
        coverage = student_exam_coverage(scoped_records, selected_student, selected_class, manual_order) if selected_student else {"present": [], "missing": []}
        return render_template(
            "index.html",
            records=records,
            lines=lines,
            classes=classes,
            selected_class=selected_class,
            exam_order=ordered_exams,
            all_exams=exam_names(records, manual_order),
            all_students=student_names(records),
            scoped_count=len(scoped_records),
            summaries=class_summary(scoped_records, lines, manual_order) if scoped_records else [],
            progress_data=progress_matrix(scoped_records, manual_order, progress_sort),
            progress_sort=progress_sort,
            students=scoped_students,
            selected_student=selected_student,
            student_coverage=coverage,
            combined_chart=combined_trend_chart_svg(series, f"{selected_student} 成绩与位次趋势", exam_order=ordered_exams, reference_lines=lines) if series else "",
        )

    @post_route(app, "/settings")
    def settings():
        lines = [
            ScoreLine("优秀", float(request.form["excellent"])),
            ScoreLine("良好", float(request.form["good"])),
            ScoreLine("及格", float(request.form["pass"])),
            ScoreLine("低分", float(request.form["low"])),
        ]
        save_lines(lines)
        flash("分数线已保存。")
        return redirect(url_for("index"))

    @post_route(app, "/import/excel")
    def import_excel():
        files = [file for file in request.files.getlist("file") if file and file.filename]
        class_name = request.form.get("class_name", "").strip()
        exam_name = request.form.get("exam_name", "").strip()
        if not files:
            flash("请选择 Excel 文件。")
            return redirect(url_for("index"))
        records: list[ScoreRecord] = []
        preview_rows: list[dict[str, object]] = []
        diagnostics: list[dict[str, object]] = []
        source_summaries: list[dict[str, object]] = []
        for file_index, file in enumerate(files, start=1):
            filename = upload_filename(file.filename, ".xlsx")
            path = UPLOAD_DIR / filename
            file.save(path)
            try:
                file_records, file_diagnostics = parse_excel_with_report(path, class_name, exam_name, file.filename)
            except InvalidFileException:
                flash(f"{file.filename} 无法读取：请确认文件是 .xlsx/.xlsm 格式，不是旧版 .xls 或图片截图。")
                return redirect(url_for("index"))
            except Exception as exc:
                flash(f"{file.filename} 导入失败：{exc}")
                return redirect(url_for("index"))
            for item in file_diagnostics:
                item["source_name"] = file.filename or filename
                item["file_index"] = file_index
            records.extend(file_records)
            source_name = file.filename or filename
            for record in file_records:
                preview_rows.append(record_to_dict(record, source_name))
            diagnostics.extend(file_diagnostics)
            source_summaries.append(
                {
                    "source_name": source_name,
                    "record_count": len(file_records),
                    "sheet_count": len(file_diagnostics),
                    "class_names": sorted({record.class_name for record in file_records}),
                    "exam_names": sorted({record.exam_name for record in file_records}),
                }
            )
        if not records:
            flash("Excel 中没有识别到成绩。请确认表头包含“姓名”，成绩列包含“语文/分数/成绩”，或使用“姓名 + 考试名称”格式。")
            return redirect(url_for("index"))
        token = uuid4().hex
        payload = {
            "records": preview_rows,
            "diagnostics": diagnostics,
            "source_name": "、".join(item["source_name"] for item in source_summaries),
            "source_summaries": source_summaries,
        }
        save_preview(token, payload)
        return render_template(
            "review_excel.html",
            token=token,
            source_name=payload["source_name"],
            source_summaries=source_summaries,
            diagnostics=diagnostics,
            batch_mappings=batch_mapping_items(preview_rows),
            preview_records=preview_rows[:12],
            record_count=len(records),
        )

    @post_route(app, "/import/excel/confirm")
    def confirm_excel():
        token = request.form.get("token", "")
        payload = load_preview(token)
        if not payload:
            flash("导入预览已失效，请重新上传 Excel。")
            return redirect(url_for("index"))
        class_map = form_mapping("class_key", "class_value")
        exam_map = form_mapping("exam_key", "exam_value")
        records = [
            ScoreRecord(
                class_name=class_map.get(row.get("class_key", ""), row["class_name"]).strip() or "未分班",
                student_name=row["student_name"],
                exam_name=exam_map.get(row.get("exam_key", ""), row["exam_name"]).strip() or "未命名考试",
                score=float(row["score"]),
                rank=int(row["rank"]) if row.get("rank") not in (None, "") else None,
            )
            for row in payload.get("records", [])
        ]
        append_scores(records)
        delete_preview(token)
        flash(f"已确认导入 {len(records)} 条 Excel 成绩。")
        return redirect(url_for("index"))

    @post_route(app, "/import/image")
    def import_image():
        files = [file for file in request.files.getlist("file") if file and file.filename]
        if not files:
            flash("请选择图片文件。")
            return redirect(url_for("index"))
        image_items: list[dict[str, object]] = []
        for file in files:
            filename = upload_filename(file.filename, ".png")
            path = UPLOAD_DIR / filename
            file.save(path)
            text, error = image_to_text(path)
            image_items.append(
                {
                    "source_name": file.filename or filename,
                    "class_name": "",
                    "exam_name": "",
                    "ocr_text": text,
                    "error": error,
                }
            )
        return render_template(
            "review_image.html",
            image_items=image_items,
        )

    @post_route(app, "/import/review")
    def import_review():
        class_names_form = request.form.getlist("class_name")
        exam_names_form = request.form.getlist("exam_name")
        table_texts = request.form.getlist("table_text")
        if not table_texts:
            table_texts = [request.form.get("table_text", "")]
        records: list[ScoreRecord] = []
        for index, table_text in enumerate(table_texts):
            class_name = (class_names_form[index] if index < len(class_names_form) else "").strip() or "未分班"
            exam_name = (exam_names_form[index] if index < len(exam_names_form) else "").strip()
            records.extend(parse_review_text(table_text, class_name, exam_name))
        append_scores(records)
        flash(f"已确认导入 {len(records)} 条图片/手工校对成绩。")
        return redirect(url_for("index"))

    @app.get("/export.xlsx")
    def export_xlsx():
        records = load_scores()
        lines = load_lines(DEFAULT_LINES)
        manual_order = load_exam_order()
        export_path = Path("data") / "语文成绩分析.xlsx"
        export_path.parent.mkdir(exist_ok=True)
        workbook = Workbook()

        detail = workbook.active
        detail.title = "成绩明细"
        detail.append(["班级", "姓名", "考试", "语文成绩", "班级位次"])
        for record in records:
            detail.append([record.class_name, record.student_name, record.exam_name, record.score, record.rank])

        summary_sheet = workbook.create_sheet("班级总览")
        summary_headers = ["班级", "考试", "人数", "均分", "中位数", "最高", "最低"] + [f"{line.label}率" for line in lines]
        summary_sheet.append(summary_headers)
        for row in class_summary(records, lines, manual_order):
            summary_sheet.append(
                [
                    row["class_name"],
                    row["exam_name"],
                    row["count"],
                    row["avg"],
                    row["median"],
                    row["max"],
                    row["min"],
                    *[row[f"{line.label}率"] for line in lines],
                ]
            )

        progress_sheet = workbook.create_sheet("进退步概览")
        progress_data = progress_matrix(records, manual_order, "name")
        progress_exams = list(progress_data["exams"])
        progress_sheet.append(["班级", "姓名", "考试次数", *progress_exams, "总分变化", "位次进步"])
        for row in progress_data["rows"]:
            progress_sheet.append(
                [
                    row["class_name"],
                    row["student_name"],
                    row["exam_count"],
                    *[format_progress_cell(cell) for cell in row["cells"]],
                    row["score_delta_total"],
                    row["rank_delta_total"],
                ]
            )

        workbook.save(export_path)
        return send_file(export_path.resolve(), as_attachment=True, download_name="语文成绩分析.xlsx")

    @app.get("/export/student-chart.svg")
    def export_student_chart():
        records = load_scores()
        lines = load_lines(DEFAULT_LINES)
        manual_order = load_exam_order()
        class_name = request.args.get("class_name", "").strip()
        student_name = request.args.get("student", "").strip()
        scoped_records = filter_by_class(records, class_name)
        series = student_series(scoped_records, student_name, class_name, manual_order) if student_name else []
        if not series:
            flash("没有找到该学生的趋势图数据。")
            return redirect(url_for("index", class_name=class_name, student=student_name))

        svg, suffix = student_chart_svg(series, scoped_records, student_name, lines, manual_order)
        return send_svg(svg, f"{safe_file_part(student_name)}_{suffix}.svg")

    @app.get("/export/student-chart.png")
    def export_student_chart_png():
        records = load_scores()
        lines = load_lines(DEFAULT_LINES)
        manual_order = load_exam_order()
        class_name = request.args.get("class_name", "").strip()
        student_name = request.args.get("student", "").strip()
        scoped_records = filter_by_class(records, class_name)
        series = student_series(scoped_records, student_name, class_name, manual_order) if student_name else []
        if not series:
            flash("没有找到该学生的趋势图数据。")
            return redirect(url_for("index", class_name=class_name, student=student_name))

        png, suffix = student_chart_png(series, scoped_records, student_name, lines, manual_order)
        return send_png(png, f"{safe_file_part(student_name)}_{suffix}.png")

    @app.get("/export/student-charts.zip")
    def export_student_charts():
        records = load_scores()
        lines = load_lines(DEFAULT_LINES)
        manual_order = load_exam_order()
        class_name = request.args.get("class_name", "").strip()
        scoped_records = filter_by_class(records, class_name)
        students = student_names(scoped_records)
        if not students:
            flash("当前班级没有可导出的学生趋势图。")
            return redirect(url_for("index", class_name=class_name))

        folder = safe_file_part(class_name) or "全部班级"
        buffer = BytesIO()
        with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
            for student_name in students:
                series = student_series(scoped_records, student_name, class_name, manual_order)
                if not series:
                    continue
                png, suffix = student_chart_png(series, scoped_records, student_name, lines, manual_order)
                student_part = safe_file_part(student_name)
                archive.writestr(f"{folder}/{student_part}_{suffix}.png", png)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype="application/zip",
            as_attachment=True,
            download_name=f"{folder}_学生折线图.zip",
        )

    @post_route(app, "/clear")
    def clear():
        save_scores([])
        flash("已清空当前成绩数据。")
        return redirect(url_for("index"))

    @post_route(app, "/delete-records")
    def delete_records():
        records = load_scores()
        mode = request.form.get("delete_mode", "").strip()
        class_name = request.form.get("class_name", "").strip()
        exam_name = request.form.get("exam_name", "").strip()
        student_name = request.form.get("student_name", "").strip()

        kept: list[ScoreRecord] = []
        removed = 0
        for record in records:
            should_remove = should_delete_record(record, mode, class_name, exam_name, student_name)
            if should_remove:
                removed += 1
            else:
                kept.append(record)
        if removed:
            save_scores(kept)
            flash(f"已删除 {removed} 条成绩记录。")
        else:
            flash("没有找到符合条件的成绩记录，未删除任何数据。")
        return redirect(url_for("index", class_name=class_name if class_name else None))

    @post_route(app, "/exam-order")
    def exam_order():
        ordered_pairs = []
        originals = request.form.getlist("exam_name")
        for exam_name in originals:
            order_text = request.form.get(f"order_{exam_name}", "").strip()
            try:
                order_value = int(order_text)
            except ValueError:
                order_value = 9999
            ordered_pairs.append((order_value, exam_name))
        save_exam_order([exam_name for _order, exam_name in sorted(ordered_pairs)])
        flash("考试顺序已保存。")
        class_name = request.form.get("class_name", "")
        return redirect(url_for("index", class_name=class_name))

    return app


def resource_path() -> Path:
    if hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)  # type: ignore[attr-defined]
    return Path(__file__).resolve().parent.parent


def post_route(app: Flask, rule: str):
    return app.route(rule, methods=["POST"])


def upload_filename(raw_name: str | None, default_suffix: str) -> str:
    raw_path = Path(raw_name or "")
    suffix = raw_path.suffix.lower() or default_suffix
    stem = secure_filename(raw_path.stem) or uuid4().hex
    return f"{stem}{suffix}"


def student_chart_svg(
    series: list[ScoreRecord],
    scoped_records: list[ScoreRecord],
    student_name: str,
    lines: list[ScoreLine],
    manual_order: list[str] | None,
) -> tuple[str, str]:
    ordered_exams = exam_order_labels(scoped_records, manual_order)
    svg = combined_trend_chart_svg(
        series,
        f"{student_name} 成绩与位次趋势",
        exam_order=ordered_exams,
        reference_lines=lines,
    )
    return svg_document(svg), "成绩位次折线图"


def student_chart_png(
    series: list[ScoreRecord],
    scoped_records: list[ScoreRecord],
    student_name: str,
    lines: list[ScoreLine],
    manual_order: list[str] | None,
) -> tuple[bytes, str]:
    ordered_exams = exam_order_labels(scoped_records, manual_order)
    png = combined_trend_chart_png(
        series,
        f"{student_name} 成绩与位次趋势",
        exam_order=ordered_exams,
        reference_lines=lines,
    )
    return png, "成绩位次折线图"


def svg_document(svg: str) -> str:
    return '<?xml version="1.0" encoding="UTF-8"?>\n' + svg.strip() + "\n"


def send_svg(svg: str, download_name: str):
    buffer = BytesIO(svg.encode("utf-8"))
    return send_file(
        buffer,
        mimetype="image/svg+xml; charset=utf-8",
        as_attachment=True,
        download_name=download_name,
    )


def send_png(png: bytes, download_name: str):
    return send_file(
        BytesIO(png),
        mimetype="image/png",
        as_attachment=True,
        download_name=download_name,
    )


def format_progress_cell(cell: dict[str, object]) -> str:
    score = cell.get("score")
    if score is None:
        return ""
    rank = cell.get("rank")
    score_text = f"{float(score):g}"
    if rank is None:
        return score_text
    return f"{score_text}（第{int(rank)}名）"


def safe_file_part(value: str) -> str:
    cleaned = "".join("_" if char in '<>:"/\\|?*' else char for char in value.strip())
    cleaned = "_".join(part for part in cleaned.split() if part)
    return cleaned.strip("._ ") or "未命名"


def record_to_dict(record: ScoreRecord, source_name: str = "") -> dict[str, object]:
    class_key = f"{source_name}||{record.class_name}" if source_name else record.class_name
    exam_key = f"{source_name}||{record.exam_name}" if source_name else record.exam_name
    return {
        "source_name": source_name,
        "class_key": class_key,
        "exam_key": exam_key,
        "class_name": record.class_name,
        "student_name": record.student_name,
        "exam_name": record.exam_name,
        "score": record.score,
        "rank": record.rank,
    }


def batch_mapping_items(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    seen: dict[tuple[str, str, str], dict[str, object]] = {}
    for row in rows:
        source_name = str(row.get("source_name", ""))
        class_name = str(row.get("class_name", ""))
        exam_name = str(row.get("exam_name", ""))
        key = (source_name, class_name, exam_name)
        item = seen.setdefault(
            key,
            {
                "source_name": source_name,
                "class_key": str(row.get("class_key", "")),
                "exam_key": str(row.get("exam_key", "")),
                "original_class": class_name,
                "original_exam": exam_name,
                "suggested_class": "" if class_name == "未分班" else class_name,
                "suggested_exam": "" if exam_name == "未命名考试" else exam_name,
                "needs_attention": class_name == "未分班" or exam_name == "未命名考试",
                "record_count": 0,
                "sample_names": [],
            },
        )
        item["record_count"] = int(item["record_count"]) + 1
        sample_names = item["sample_names"]
        if isinstance(sample_names, list) and len(sample_names) < 3:
            sample_names.append(str(row.get("student_name", "")))
    return sorted(
        seen.values(),
        key=lambda item: (str(item["source_name"]), str(item["original_class"]), str(item["original_exam"])),
    )


def mapping_items(values: set[str]) -> list[dict[str, object]]:
    return [
        {
            "original": value,
            "suggested": "" if value in ("未分班", "未命名考试") else value,
            "needs_attention": value in ("未分班", "未命名考试"),
        }
        for value in sorted(values)
    ]


def form_mapping(original_name: str, value_name: str) -> dict[str, str]:
    originals = request.form.getlist(original_name)
    values = request.form.getlist(value_name)
    return dict(zip(originals, values))


def should_delete_record(
    record: ScoreRecord,
    mode: str,
    class_name: str,
    exam_name: str,
    student_name: str,
) -> bool:
    if mode == "class":
        return bool(class_name) and record.class_name == class_name
    if mode == "exam":
        return bool(exam_name) and record.exam_name == exam_name
    if mode == "student":
        return bool(student_name) and record.student_name == student_name
    if mode == "class_exam":
        return bool(class_name and exam_name) and record.class_name == class_name and record.exam_name == exam_name
    if mode == "class_student":
        return bool(class_name and student_name) and record.class_name == class_name and record.student_name == student_name
    return False


def parse_review_text(text: str, class_name: str, exam_name: str) -> list[ScoreRecord]:
    table_records = parse_review_table(text, class_name, exam_name)
    if table_records:
        return table_records

    records: list[ScoreRecord] = []
    for line in text.splitlines():
        parts = [part.strip() for part in line.replace("\t", " ").split(" ") if part.strip()]
        if len(parts) < 2:
            continue
        name = parts[0]
        score = coerce_score(parts[1])
        if not name or score is None or name in {"姓名", "语文"}:
            continue
        rank = int(parts[2]) if len(parts) > 2 and parts[2].isdigit() else None
        records.append(ScoreRecord(class_name, name, exam_name, score, rank))
    return competition_ranks(records) if any(record.rank is None for record in records) else records


def parse_review_table(text: str, class_name: str, exam_name: str) -> list[ScoreRecord]:
    rows = [split_review_row(line) for line in text.splitlines()]
    rows = [row for row in rows if row]
    header_index = next((index for index, row in enumerate(rows) if any("姓名" in item for item in row)), None)
    if header_index is None:
        return []

    headers = rows[header_index]
    name_col = next((index for index, item in enumerate(headers) if "姓名" in item), None)
    if name_col is None:
        return []
    rank_col = next(
        (index for index, item in enumerate(headers) if any(key in item for key in ("排名", "位次", "名次"))),
        None,
    )

    score_cols: list[tuple[int, str]] = []
    for index, header in enumerate(headers):
        if index == name_col or index == rank_col:
            continue
        if any(key in header for key in ("语文", "分数", "成绩")) or ("期" in header and not any(key in header for key in ("排名", "位次", "名次"))):
            cleaned = clean_review_exam_name(header)
            score_cols.append((index, exam_name if cleaned == "语文" and exam_name else cleaned))
    if not score_cols:
        return []

    records: list[ScoreRecord] = []
    for row in rows[header_index + 1 :]:
        if name_col >= len(row):
            continue
        name = row[name_col].strip()
        if not name:
            continue
        for col, resolved_exam in score_cols:
            if col >= len(row):
                continue
            score = coerce_score(row[col])
            if score is None:
                continue
            rank = None
            if rank_col is not None and rank_col < len(row) and len(score_cols) == 1:
                rank = parse_rank_text(row[rank_col])
            records.append(ScoreRecord(class_name, name, resolved_exam or exam_name, score, rank))
    return competition_ranks(records) if any(record.rank is None for record in records) else records


def split_review_row(line: str) -> list[str]:
    if "\t" in line:
        return [part.strip() for part in line.split("\t") if part.strip()]
    return [part.strip() for part in line.replace("　", " ").split(" ") if part.strip()]


def clean_review_exam_name(header: str) -> str:
    cleaned = header.strip()
    for token in ("语文班级排名", "班级排名", "语文", "分数", "成绩", "排名", "位次", "名次"):
        cleaned = cleaned.replace(token, "")
    return cleaned.strip() or header.strip()


def parse_rank_text(value: str) -> int | None:
    digits = "".join(ch for ch in value if ch.isdigit())
    return int(digits) if digits else None
