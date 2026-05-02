from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from .models import ScoreRecord, coerce_score, competition_ranks


NAME_KEYS = ("姓名", "学生", "名字")
SCORE_KEYS = ("语文", "分数", "成绩")
RANK_KEYS = ("排名", "位次", "名次")
NON_SCORE_KEYS = ("学号", "序号", "编号", "班级", "姓名", "学生", "名字", "排名", "位次", "名次")
GENERIC_EXAM_NAMES = ("语文", "分数", "成绩", "语文综合排名表", "Sheet", "Sheet1", "未命名考试")


def parse_excel(path: Path, class_name: str, fallback_exam_name: str, source_name: str | None = None) -> list[ScoreRecord]:
    records, _diagnostics = parse_excel_with_report(path, class_name, fallback_exam_name, source_name)
    return records


def parse_excel_with_report(
    path: Path,
    class_name: str,
    fallback_exam_name: str,
    source_name: str | None = None,
) -> tuple[list[ScoreRecord], list[dict[str, object]]]:
    workbook = load_workbook(BytesIO(path.read_bytes()), data_only=True)
    records: list[ScoreRecord] = []
    diagnostics: list[dict[str, object]] = []
    source_text = source_name or path.name
    source_class_name = infer_class_name(source_text)
    source_exam_name = infer_exam_name(source_text)
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        sheet_records, diagnostic = _parse_sheet(
            rows,
            class_name,
            fallback_exam_name,
            sheet.title,
            source_class_name,
            source_exam_name,
        )
        records.extend(sheet_records)
        diagnostics.append(diagnostic)
    return _fill_missing_ranks(records), diagnostics


def _parse_sheet(
    rows: list[tuple[object, ...]],
    class_name: str,
    fallback_exam_name: str,
    sheet_title: str,
    source_class_name: str | None,
    source_exam_name: str | None,
) -> tuple[list[ScoreRecord], dict[str, object]]:
    header_index = _find_header_row(rows)
    if header_index is None:
        return [], {
            "sheet_title": sheet_title,
            "class_name": "",
            "class_source": "未读取到",
            "exam_names": [],
            "record_count": 0,
            "message": "没有找到包含“姓名”的表头行。",
        }

    headers = [str(value).strip() if value is not None else "" for value in rows[header_index]]
    name_col = _find_col(headers, NAME_KEYS)
    if name_col is None:
        return [], {
            "sheet_title": sheet_title,
            "class_name": "",
            "class_source": "未读取到",
            "exam_names": [],
            "record_count": 0,
            "message": "没有找到姓名列。",
        }

    sheet_class_name, class_source = _sheet_class_name(
        rows,
        header_index,
        class_name,
        sheet_title,
        source_class_name,
    )
    sheet_fallback_exam = fallback_exam_name or source_exam_name or infer_exam_name(sheet_title) or "未命名考试"
    rank_col = _find_col(headers, RANK_KEYS)
    score_cols = _score_columns(headers, name_col, rank_col, sheet_fallback_exam)
    records: list[ScoreRecord] = []

    for row in rows[header_index + 1 :]:
        if name_col >= len(row):
            continue
        student_name = str(row[name_col]).strip() if row[name_col] is not None else ""
        if not student_name:
            continue
        for col, exam_name in score_cols:
            if col >= len(row):
                continue
            score = coerce_score(row[col])
            if score is None:
                continue
            rank = _parse_rank(row[rank_col]) if rank_col is not None and len(row) > rank_col and len(score_cols) == 1 else None
            records.append(
                ScoreRecord(
                    class_name=sheet_class_name,
                    student_name=student_name,
                    exam_name=exam_name or fallback_exam_name,
                    score=score,
                    rank=rank,
                )
            )
    exam_names = sorted({record.exam_name for record in records})
    return records, {
        "sheet_title": sheet_title,
        "class_name": sheet_class_name,
        "class_source": class_source,
        "exam_names": exam_names,
        "record_count": len(records),
        "message": "",
    }


def _find_header_row(rows: list[tuple[object, ...]]) -> int | None:
    for index, row in enumerate(rows[:10]):
        values = [str(value).strip() for value in row if value is not None]
        if any(any(key in value for key in NAME_KEYS) for value in values):
            return index
    return None


def _find_col(headers: list[str], keys: tuple[str, ...]) -> int | None:
    for index, header in enumerate(headers):
        if any(key in header for key in keys):
            return index
    return None


def _score_columns(headers: list[str], name_col: int, rank_col: int | None, fallback_exam_name: str) -> list[tuple[int, str]]:
    explicit: list[tuple[int, str]] = []
    for index, header in enumerate(headers):
        if index == name_col or index == rank_col:
            continue
        if any(key in header for key in SCORE_KEYS) and not any(key in header for key in RANK_KEYS):
            explicit.append((index, _clean_exam_name(header, fallback_exam_name)))
    if explicit:
        return explicit

    inferred: list[tuple[int, str]] = []
    for index, header in enumerate(headers):
        if index == name_col or index == rank_col:
            continue
        if header and not any(key in header for key in NON_SCORE_KEYS):
            inferred.append((index, _clean_exam_name(header, fallback_exam_name)))
    return inferred


def _clean_exam_name(header: str, fallback_exam_name: str) -> str:
    cleaned = re.sub(r"(语文|分数|成绩|班级|排名|位次|名次)", "", header).strip()
    if cleaned:
        return cleaned
    if header.strip() in GENERIC_EXAM_NAMES and fallback_exam_name:
        return fallback_exam_name
    return header.strip() or fallback_exam_name


def _sheet_class_name(
    rows: list[tuple[object, ...]],
    header_index: int,
    fallback_class_name: str,
    sheet_title: str,
    source_class_name: str | None,
) -> tuple[str, str]:
    for row in rows[:header_index]:
        for value in row:
            if value is None:
                continue
            text = str(value).strip()
            if "班" in text:
                return normalize_class_name(text), "表格内容"
    sheet_class_name = infer_class_name(sheet_title)
    if sheet_class_name:
        return sheet_class_name, "工作表名"
    if source_class_name:
        return source_class_name, "文件名"
    if fallback_class_name:
        return normalize_class_name(fallback_class_name), "手动兜底"
    return "未分班", "未读取到"


def infer_class_name(text: str) -> str | None:
    normalized = normalize_class_name(text)
    match = re.search(r"([0-9]+)班", normalized)
    if match:
        return f"{match.group(1)}班"
    return None


def normalize_class_name(text: str) -> str:
    chinese_digits = {
        "一": "1",
        "二": "2",
        "两": "2",
        "三": "3",
        "四": "4",
        "五": "5",
        "六": "6",
        "七": "7",
        "八": "8",
        "九": "9",
        "十": "10",
    }
    normalized = str(text).strip()
    for chinese, digit in chinese_digits.items():
        normalized = re.sub(fr"{chinese}(?=班)", digit, normalized)
    return normalized


def infer_exam_name(text: str) -> str | None:
    compact = re.sub(r"\s+", "", str(text))
    grade_semester = r"(七上|七下|八上|八下|九上|九下)"
    phase = r"(期中|期末|第一次月考|第二次月考|第三次月考|第1次月考|第2次月考|第3次月考)"
    match = re.search(grade_semester + r".*?" + phase, compact)
    if match:
        return f"{match.group(1)}{match.group(2)}"
    match = re.search(phase, compact)
    if match:
        return match.group(1)
    return None


def _parse_rank(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r"\d+", str(value))
    return int(match.group(0)) if match else None


def _fill_missing_ranks(records: list[ScoreRecord]) -> list[ScoreRecord]:
    missing = [record for record in records if record.rank is None]
    if not missing:
        return records
    ranked_missing = competition_ranks(missing)
    existing = [record for record in records if record.rank is not None]
    return existing + ranked_missing
