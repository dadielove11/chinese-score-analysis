from __future__ import annotations

from io import BytesIO
from pathlib import Path
import os
import sys
import tempfile
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from student_analysis.charts import combined_trend_chart_svg
import student_analysis.charts as charts
from student_analysis.models import DEFAULT_LINES, ScoreRecord
from student_analysis.storage import save_scores
from student_analysis.web import create_app


def sample_records() -> list[ScoreRecord]:
    return [
        ScoreRecord("7班", "张三", "七上期中", 72.5, 18),
        ScoreRecord("7班", "张三", "七上期末", 81, 8),
        ScoreRecord("7班", "张三", "七下期中", 78, 11),
        ScoreRecord("7班", "李四", "七上期中", 66, 25),
        ScoreRecord("7班", "李四", "七上期末", 70, 20),
    ]


def test_default_score_lines_are_teacher_thresholds() -> None:
    assert [(line.label, line.score) for line in DEFAULT_LINES] == [
        ("优秀", 80),
        ("良好", 70),
        ("及格", 60),
        ("低分", 40),
    ]


def test_combined_chart_legend_uses_colored_lines() -> None:
    svg = combined_trend_chart_svg(
        sample_records()[:3],
        "张三 成绩与位次趋势",
        exam_order=["七上期中", "七上期末", "七下期中"],
    )

    assert "score-legend-line" in svg
    assert "rank-legend-line" in svg
    assert "legend-point" not in svg


def test_export_progress_sheet_expands_each_exam_column() -> None:
    root = Path.cwd()
    tmp = tempfile.TemporaryDirectory()
    try:
        os.chdir(tmp.name)
        save_scores(sample_records())
        response = create_app().test_client().get("/export.xlsx")
        assert response.status_code == 200

        workbook = load_workbook(BytesIO(response.data), data_only=True)
        response.close()
        sheet = workbook["进退步概览"]
        headers = [cell.value for cell in sheet[1]]
        assert headers == [
            "班级",
            "姓名",
            "考试次数",
            "七上期中",
            "七上期末",
            "七下期中",
            "总分变化",
            "位次进步",
        ]
        first_row = [cell.value for cell in sheet[2]]
        assert first_row[0:3] == ["7班", "张三", 3]
        assert first_row[3:6] == ["72.5（第18名）", "81（第8名）", "78（第11名）"]
        workbook.close()
    finally:
        os.chdir(root)
        tmp.cleanup()


def test_template_download_contains_supported_layouts() -> None:
    response = create_app().test_client().get("/template.xlsx")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    workbook = load_workbook(BytesIO(response.data), data_only=True)
    assert workbook.sheetnames == ["单次考试模板", "多次考试模板"]
    assert [cell.value for cell in workbook["单次考试模板"][1]] == ["班级", "姓名", "考试", "语文", "位次"]
    assert [cell.value for cell in workbook["多次考试模板"][1]][:4] == ["班级", "姓名", "七上期中", "七上期末"]
    workbook.close()


def test_index_uses_single_import_entry() -> None:
    response = create_app().test_client().get("/")
    assert response.status_code == 200

    html = response.get_data(as_text=True)
    assert "导入成绩文件" in html
    assert "开始导入" in html
    assert "下载 Excel 模板" in html
    assert 'action="/import/files"' in html
    assert "Excel 导入" not in html
    assert "图片导入" not in html


def test_unified_import_dispatches_excel_file() -> None:
    root = Path.cwd()
    tmp = tempfile.TemporaryDirectory()
    try:
        os.chdir(tmp.name)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "七上期中"
        sheet.append(["班级", "姓名", "语文", "位次"])
        sheet.append(["7班", "张三", 88, 1])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)

        response = create_app().test_client().post(
            "/import/files",
            data={"file": (payload, "成绩.xlsx")},
            content_type="multipart/form-data",
        )

        assert response.status_code == 200
        assert "导入批次确认" in response.get_data(as_text=True)
    finally:
        os.chdir(root)
        tmp.cleanup()


def test_unified_import_dispatches_image_file() -> None:
    root = Path.cwd()
    tmp = tempfile.TemporaryDirectory()
    try:
        os.chdir(tmp.name)
        image = Image.new("RGB", (40, 40), "white")
        payload = BytesIO()
        image.save(payload, format="PNG")
        payload.seek(0)

        with patch("student_analysis.web.image_to_text", return_value=("张三 88 1", None)) as ocr:
            response = create_app().test_client().post(
                "/import/files",
                data={"file": (payload, "截图.png")},
                content_type="multipart/form-data",
            )

        html = response.get_data(as_text=True)
        assert response.status_code == 200
        assert "图片识别校对" in html
        assert "张三 88 1" in html
        assert ocr.called
    finally:
        os.chdir(root)
        tmp.cleanup()


def test_combined_chart_png_rotates_crowded_axis_labels() -> None:
    records = [
        ScoreRecord("7班", "张三", f"七上第{index}次阶段检测", 60 + index, 30 - index)
        for index in range(1, 13)
    ]
    exam_order = [record.exam_name for record in records]

    with patch.object(charts, "draw_rotated_centered_text") as rotated_text:
        png = charts.combined_trend_chart_png(records, "张三 成绩与位次趋势", exam_order=exam_order)

    label_calls = [call for call in rotated_text.call_args_list if call.args and call.args[1] in exam_order]
    assert len(label_calls) == len(exam_order)
    assert min(call.args[2] for call in label_calls) >= 120
    assert Image.open(BytesIO(png)).height >= 1000


if __name__ == "__main__":
    test_default_score_lines_are_teacher_thresholds()
    test_combined_chart_legend_uses_colored_lines()
    test_export_progress_sheet_expands_each_exam_column()
    test_template_download_contains_supported_layouts()
    test_index_uses_single_import_entry()
    test_unified_import_dispatches_excel_file()
    test_unified_import_dispatches_image_file()
    test_combined_chart_png_rotates_crowded_axis_labels()
